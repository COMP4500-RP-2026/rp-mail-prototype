"""Personal RP+ mail workspace for Windows and macOS."""
import csv
import io
import json
import os
from pathlib import Path
import queue
import sqlite3
import sys
import threading
import traceback
import webbrowser
from contextlib import closing, redirect_stdout
from types import SimpleNamespace
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import rp_mailer as core
import email_lookup
from localization import tr, set_language, DisplayVar

UI_FONT = 'Helvetica Neue' if sys.platform == 'darwin' else 'Microsoft YaHei UI'
PREVIEW_FONT = 'Helvetica Neue' if sys.platform == 'darwin' else 'Segoe UI'

def app_home():
    if '--self-check' in sys.argv and sys.argv.index('--self-check') + 1 < len(sys.argv):
        return Path(sys.argv[sys.argv.index('--self-check') + 1]).resolve().parent / 'rpmail-self-check-data'
    if getattr(sys, 'frozen', False):
        if sys.platform == 'darwin':
            return Path.home() / 'Library' / 'Application Support' / 'RPMail'
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent

HOME = app_home()
core.BASE = HOME
DATA = HOME / 'personal_data'
LABELS = {
    'researcher_name': '研究人员姓名', 'researcher_email': '收件邮箱',
    'title': '研究成果标题', 'project_name': '项目名称', 'project_code': '资助编号',
    'publication_date': '发表日期', 'funder': '资助机构',
    'project_start_date': '项目开始日期', 'project_end_date': '项目结束日期',
    'output_url': 'RP+ 成果链接', 'suggestion_reason': '推荐理由（英文，需核实）',
}
HISTORY = {'accepted': '服务器已接受', 'submitted_outlook': '已交给 Outlook',
           'drafted': '已存草稿', 'pending': '结果待核查', 'uncertain': '结果待核查'}

def load_table(path):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = wb.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(x).strip() if x is not None else '' for x in next(iterator)]
            rows = [{k: str(v).strip() if v is not None else '' for k, v in zip(headers, values) if k}
                    for values in iterator if any(v is not None for v in values)]
            info = f'已读取工作表：{sheet.title}'
        finally:
            wb.close()
    else:
        try:
            headers, rows = core.read_rows(path)
            info = '已按 UTF-8 读取'
        except UnicodeDecodeError:
            headers, rows = core.read_rows(path, 'cp1252')
            info = '源文件不是 UTF-8，已按 Windows-1252 读取；请核对标题乱码'
    required = {'output_uuid', 'project_uuid', 'title', 'output_url', 'project_code', 'already_in_relations'}
    if not required <= set(headers):
        raise ValueError('缺少原始字段：' + ', '.join(sorted(required-set(headers))))
    fields = list(dict.fromkeys([x for x in headers if x] + core.EXTRA + email_lookup.FIELDS))
    return fields, rows, info

def save_table(path, fields, rows):
    path = Path(path)
    temporary = path.with_suffix('.tmp')
    with temporary.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.ui_scale=max(1.0,float(self.tk.call('tk','scaling'))/(96/72))
        self.title(tr('RP Mail · 个人研究确认'))
        width=min(self.px(1380),int(self.winfo_screenwidth()*.94))
        height=min(self.px(880),int(self.winfo_screenheight()*.92))
        self.geometry(f'{width}x{height}')
        self.minsize(min(self.px(1200),width), min(self.px(740),height))
        self.configure(bg='#f3f6fa')
        self.rows, self.fields, self.current = [], [], None
        self.history, self.busy, self.dirty = {}, False, False
        self.events = queue.Queue()
        self.lookup_stop = threading.Event()
        self.auth_stop = threading.Event()
        self.auth_pending = threading.Event()
        self.auth_window = None
        self.close_when_idle = False
        self.filter = '全部记录'
        self.vars = {}
        self.settings = dict(transport='outlook' if sys.platform == 'win32' else 'smtp',
                             sender_name='', sender_email='', reply_to='', delay_seconds=2,
                             smtp_host='smtp.office365.com', smtp_port='587', smtp_security='starttls',
                             smtp_username='', auth_mode='microsoft', tenant_id='', client_id='',
                             password_env='RP_MAIL_SMTP_PASSWORD', language='en')
        DATA.mkdir(parents=True, exist_ok=True)
        config = DATA / 'settings.json'
        if config.exists():
            self.settings.update(json.loads(config.read_text(encoding='utf-8')))
        if sys.platform == 'darwin' and self.settings.get('transport') == 'outlook':
            self.settings['transport'] = 'smtp'
        set_language(self.settings.get('language','en'))
        self.title(tr('RP Mail · 个人研究确认'))
        self.style_ui()
        self.build()
        self.protocol('WM_DELETE_WINDOW', self.close_app)
        self.after(100, self.poll)
        saved = DATA / 'mailing_list.csv'
        bundled = HOME / 'mailing_list.csv'
        if saved.exists() or bundled.exists():
            self.open_path(saved if saved.exists() else bundled)

    def style_ui(self):
        s = ttk.Style(self)
        s.theme_use('clam')
        s.configure('.', font=(UI_FONT, 10), background='#f3f6fa', foreground='#253248')
        s.configure('TButton', padding=(12, 7), background='#ffffff', borderwidth=0)
        s.map('TButton', background=[('active', '#e5efff')])
        s.configure('Primary.TButton', background='#0f6cbd', foreground='white')
        s.map('Primary.TButton', background=[('active', '#115ea3')])
        s.configure('Treeview', background='white', fieldbackground='white', rowheight=self.px(40), borderwidth=0)
        s.configure('Treeview.Heading', background='#edf2f7', padding=8, font=(UI_FONT, 10))
        s.map('Treeview', background=[('selected', '#dcecff')], foreground=[('selected', '#153c68')])
        s.configure('TNotebook', background='white', borderwidth=0)
        s.configure('TNotebook.Tab', padding=(20, 10))

    def build(self):
        bar = tk.Frame(self, bg='#0f6cbd', height=self.px(62))
        bar.pack(fill='x'); bar.pack_propagate(False)
        tk.Label(bar, text=tr('RP Mail'), bg='#0f6cbd', fg='white', font=(PREVIEW_FONT, 21, 'bold')).pack(side='left', padx=24)
        tk.Label(bar, text=tr('个人研究确认'), bg='#0f6cbd', fg='#e3efff', font=(UI_FONT, 11)).pack(side='left')
        self.account_label = tk.Label(bar, text=tr('尚未设置个人邮箱'), bg='#0f6cbd', fg='white')
        self.account_label.pack(side='right', padx=22)
        self.language_choice=tk.StringVar(value='中文' if self.settings.get('language')=='zh' else 'English')
        self.language_selector=ttk.Combobox(bar,textvariable=self.language_choice,values=('English','中文'),state='readonly',width=10)
        self.language_selector.pack(side='right',padx=10)
        self.language_selector.bind('<<ComboboxSelected>>',self.switch_language)
        toolbar = ttk.Frame(self, padding=(16, 12))
        toolbar.pack(fill='x')
        for title, command, primary in [('导入表格', self.import_file, True), ('自动查找邮箱', self.find_emails, True), ('导出清单', self.export, False),
                ('个人邮箱设置', self.account_settings, False), ('检查发件连接', self.check_connection, False)]:
            ttk.Button(toolbar, text=tr(title), command=command, style='Primary.TButton' if primary else 'TButton').pack(side='left', padx=4)
        self.stop_button=ttk.Button(toolbar,text=tr('停止查找'),command=self.lookup_stop.set,state='disabled')
        self.stop_button.pack(side='left',padx=4)
        self.search = tk.StringVar()
        ttk.Entry(toolbar, textvariable=self.search, width=20).pack(side='right', padx=10)
        ttk.Label(toolbar, text=tr('搜索')).pack(side='right')
        self.search.trace_add('write', lambda *_: self.refresh())
        body = tk.Frame(self, bg='#f3f6fa'); body.pack(fill='both', expand=True)
        sidebar = tk.Frame(body, bg='#edf2f8', width=self.px(174))
        sidebar.pack(side='left', fill='y'); sidebar.pack_propagate(False)
        tk.Label(sidebar, text=tr('邮件工作区'), bg='#edf2f8', fg='#66768b', anchor='w').pack(fill='x', padx=18, pady=(24,14))
        for name in ['全部记录','待补充','可发送','已存草稿','已提交','结果待核查']:
            tk.Button(sidebar, text=tr(name), anchor='w', bg='#edf2f8', fg='#253248', activebackground='#dcecff', relief='flat',
                      font=(UI_FONT, 11), padx=20, pady=12, command=lambda n=name:self.set_filter(n)).pack(fill='x')
        tk.Label(sidebar, text=tr('以个人身份联系\n使用自己的邮箱\n回复在邮箱中查看'), bg='#edf2f8', fg='#66768b', justify='left', font=(UI_FONT, 9)).pack(side='bottom', padx=12, pady=22)
        pane = ttk.Panedwindow(body, orient='horizontal'); pane.pack(fill='both', expand=True, padx=(0,14))
        left = ttk.Frame(pane, padding=12); right = ttk.Frame(pane, padding=12)
        pane.add(left, weight=3); pane.add(right, weight=4)
        self.count_label = ttk.Label(left, text=tr('研究关系'), font=(UI_FONT,14,'bold'))
        self.count_label.pack(anchor='w', pady=(0,12))
        list_frame = ttk.Frame(left); list_frame.pack(fill='both', expand=True)
        self.tree = ttk.Treeview(list_frame, columns=('name','title','status'), show='headings', selectmode='extended')
        for column, title, width in [('name','研究人员',110),('title','研究成果',280),('status','状态',130)]:
            self.tree.heading(column,text=tr(title)); self.tree.column(column,width=self.px(width),minwidth=self.px(80))
        scroll = ttk.Scrollbar(list_frame, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.pack(side='left',fill='both',expand=True); scroll.pack(side='right',fill='y')
        self.tree.bind('<<TreeviewSelect>>',self.select)
        selection_hint = '⌘ / Shift 多选记录；发送仅处理选中项。' if sys.platform == 'darwin' else 'Ctrl / Shift 多选记录；发送仅处理选中项。'
        ttk.Label(left,text=tr(selection_hint),foreground='#64748b').pack(anchor='w',pady=10)
        actions=ttk.Frame(left);actions.pack(fill='x')
        self.action_buttons=[]
        for title, command in [('试发给自己',lambda:self.dispatch('test')),('生成草稿',lambda:self.dispatch('draft')),('发送选中项',lambda:self.dispatch('send'))]:
            b=ttk.Button(actions,text=tr(title),command=command,style='Primary.TButton' if '发送' in title else 'TButton');b.pack(side='left',padx=3);self.action_buttons.append(b)
            if title == '生成草稿':self.draft_button=b
        tabs=ttk.Notebook(right);tabs.pack(fill='both',expand=True)
        edit=ttk.Frame(tabs,padding=14); view=ttk.Frame(tabs,padding=14); contacts=ttk.Frame(tabs,padding=14)
        tabs.add(edit,text=tr('收件人和研究信息'));tabs.add(view,text=tr('邮件预览'))
        tabs.add(contacts,text=tr('邮箱来源与候选'))
        ttk.Label(contacts,text=tr('项目网站上的公开联系人'),font=(UI_FONT,13,'bold')).pack(anchor='w',pady=(0,10))
        ttk.Label(contacts,text=tr('优先同一论文作者，再考虑项目负责人和其他项目成员。\n仅表示公开联系信息；研究关系仍需核实。'),wraplength=480).pack(anchor='w',pady=(0,12))
        self.contact_tree=ttk.Treeview(contacts,columns=('name','email','basis'),show='headings',height=8,selectmode='browse')
        for k,label,width in [('name','姓名',120),('email','公开邮箱',200),('basis','来源依据',170)]:
            self.contact_tree.heading(k,text=tr(label));self.contact_tree.column(k,width=self.px(width),minwidth=self.px(80))
        self.contact_tree.pack(fill='x')
        ttk.Button(contacts,text=tr('使用选中的联系人'),command=self.use_contact,style='Primary.TButton').pack(anchor='w',pady=12)
        self.source_info=DisplayVar(value='选择一条记录，点击“自动查找邮箱”。')
        ttk.Label(contacts,textvariable=self.source_info,wraplength=490,justify='left').pack(anchor='w',fill='x',pady=10)
        links=ttk.Frame(contacts);links.pack(anchor='w',pady=10)
        ttk.Button(links,text=tr('打开项目网站'),command=lambda:self.open_source('project_url')).pack(side='left',padx=3)
        ttk.Button(links,text=tr('打开邮箱来源'),command=lambda:self.open_source('email_source_url')).pack(side='left',padx=3)
        self.tabs=tabs
        for n,(field,label) in enumerate(LABELS.items()):
            ttk.Label(edit,text=tr(label)).grid(row=n,column=0,sticky='w',pady=6,padx=(0,12))
            value=tk.StringVar();self.vars[field]=value
            ttk.Entry(edit,textvariable=value).grid(row=n,column=1,sticky='ew',pady=6)
            value.trace_add('write',lambda *_:self.mark_dirty())
        edit.columnconfigure(1,weight=1)
        self.reviewed=tk.BooleanVar()
        ttk.Checkbutton(edit,text=tr('我已核对收件人、成果和推荐理由'),variable=self.reviewed,command=self.mark_dirty).grid(row=len(LABELS),column=0,columnspan=2,sticky='w',pady=12)
        ttk.Button(edit,text=tr('保存修改并更新预览'),command=self.save_current,style='Primary.TButton').grid(row=len(LABELS)+1,column=0,columnspan=2,sticky='w',pady=8)
        self.detail_status=ttk.Label(edit,text=tr('选择左侧记录开始编辑'),wraplength=460,foreground='#64748b')
        self.detail_status.grid(row=len(LABELS)+2,column=0,columnspan=2,sticky='w',pady=12)
        self.edit_controls=[w for w in edit.winfo_children() if isinstance(w,(ttk.Entry,ttk.Checkbutton,ttk.Button))]
        self.preview=tk.Text(view,wrap='word',font=(PREVIEW_FONT,11),relief='flat',padx=16,pady=16,bg='white',fg='#263448')
        vs=ttk.Scrollbar(view,command=self.preview.yview);self.preview.configure(yscrollcommand=vs.set)
        vs.pack(side='right',fill='y');self.preview.pack(fill='both',expand=True)
        self.status=DisplayVar(value='就绪 · 尚未发送任何邮件')
        ttk.Label(self,textvariable=self.status,padding=(18,10),foreground='#526479').pack(fill='x')
        self.update_account_label()

    def mark_dirty(self):
        self.dirty=True

    def px(self,value):
        return round(value*self.ui_scale)

    def switch_language(self,event=None):
        chosen='zh' if self.language_choice.get()=='中文' else 'en'
        if self.busy:
            self.language_choice.set('中文' if self.settings.get('language')=='zh' else 'English')
            messagebox.showinfo(tr('正在处理'),tr('操作结束后才能切换语言。'))
            return
        self.save_current(refresh=False)
        selection=self.tree.selection()
        index=self.current
        tab_index=self.tabs.index(self.tabs.select())
        query=self.search.get()
        status=self.status.raw
        self.settings['language']=chosen
        (DATA/'settings.json').write_text(json.dumps(self.settings,ensure_ascii=False,indent=2),encoding='utf-8')
        set_language(chosen)
        self.title(tr('RP Mail · 个人研究确认'))
        for child in self.winfo_children():child.destroy()
        self.vars={};self.current=None;self.dirty=False
        self.build()
        self.search.set(query)
        self.refresh()
        visible=[item for item in selection if self.tree.exists(item)]
        if visible:self.tree.selection_set(visible)
        if index is not None:
            self.current=index
            for key,var in self.vars.items():var.set(self.rows[index].get(key,''))
            self.reviewed.set(self.rows[index].get('reviewed','').lower()=='yes')
            self.dirty=False
            self.render_preview()
        self.tabs.select(tab_index)
        self.status.set(status)

    def update_account_label(self):
        self.account_label.config(text=tr(self.settings.get('sender_email') or '尚未设置个人邮箱'))
        self.draft_button.config(state='normal' if self.settings.get('transport') == 'outlook' else 'disabled')

    def read_history(self):
        self.history={}
        path=HOME/'send_history.sqlite3'
        if path.exists():
            with closing(sqlite3.connect(path)) as db:
                if db.execute("SELECT name FROM sqlite_master WHERE name='history'").fetchone():
                    self.history=dict(db.execute('SELECT id,status FROM history'))

    def row_status(self,row):
        if row.get('output_uuid') and row.get('project_uuid') and row.get('researcher_email'):
            prior=self.history.get(core.key_for(row))
            if prior:return HISTORY.get(prior,prior)
        if row.get('already_in_relations','').lower() in ('true','1','yes'):return '已有关系'
        return '待补充' if core.validate(row) else '可发送'

    def set_filter(self,name):
        if self.busy:return
        self.save_current(refresh=False)
        self.filter=name;self.refresh()

    def refresh(self):
        if self.busy:return
        selected=set(self.tree.selection())
        self.read_history()
        self.tree.delete(*self.tree.get_children())
        query=self.search.get().strip().lower();count=0
        for i,row in enumerate(self.rows):
            status=self.row_status(row)
            matches=self.filter=='全部记录' or status==self.filter or (self.filter=='已提交' and status in ('服务器已接受','已交给 Outlook'))
            if not matches or (query and query not in ' '.join(row.values()).lower()):continue
            self.tree.insert('', 'end',iid=str(i),values=(row.get('researcher_name') or tr('未填写'),row.get('title',''),tr(status)));count+=1
        kept=[i for i in selected if self.tree.exists(i)]
        if kept:self.tree.selection_set(kept)
        self.count_label.config(text=tr(f'{self.filter}  ·  {count:,}'))

    def select(self,event=None):
        if self.busy:return
        selected=self.tree.selection()
        if not selected:return
        index=int(selected[0])
        if index==self.current:return
        self.save_current(refresh=False)
        self.current=index
        row=self.rows[index]
        for key,value in self.vars.items():value.set(row.get(key,''))
        self.reviewed.set(row.get('reviewed','').lower()=='yes')
        self.dirty=False
        self.render_preview()

    def save_current(self,refresh=True):
        if self.busy:return
        if self.current is not None and self.dirty:
            row=self.rows[self.current]
            row.update({k:v.get().strip() for k,v in self.vars.items()})
            row['reviewed']='yes' if self.reviewed.get() else ''
            save_table(DATA/'mailing_list.csv',self.fields,self.rows)
            self.dirty=False
        self.render_preview()
        if refresh:self.refresh()

    def render_preview(self):
        if self.current is None:return
        row=self.rows[self.current]
        config=dict(self.settings)
        config['sender_name']=config.get('sender_name') or '[Your name]'
        config['sender_email']=config.get('sender_email') or 'your-email@example.com'
        complete={k:'' for k in self.fields};complete.update(row)
        msg=core.message(complete,config)
        text=f'From: {msg["From"]}\nTo: {msg["To"]}\nSubject: {msg["Subject"]}\n\n{msg.get_content()}'
        self.preview.config(state='normal');self.preview.delete('1.0','end');self.preview.insert('1.0',text);self.preview.config(state='disabled')
        reason=core.validate(row)
        self.detail_status.config(text=tr(f'{self.row_status(row)}'+(f' · {reason}' if reason else ' · 此条目已具备发送信息')))
        self.contact_tree.delete(*self.contact_tree.get_children())
        try:candidates=json.loads(row.get('email_candidates') or '[]')
        except (ValueError,TypeError):candidates=[]
        for i,c in enumerate(candidates):
            self.contact_tree.insert('','end',iid=str(i),values=(c.get('name',''),c.get('email',''),tr(c.get('basis',''))))
        self.source_info.set('\n\n'.join(x for x in [row.get('email_lookup_status','尚未查找'),
            row.get('contact_basis',''),row.get('email_source_url',''),row.get('email_lookup_time',''),row.get('email_lookup_warnings','')] if x))

    def open_source(self,field):
        if self.current is None:return
        url=self.rows[self.current].get(field,'')
        if not url:messagebox.showinfo(tr('暂无来源'),tr('该条记录尚无此来源链接。'));return
        try:url=email_lookup.canonical(url)
        except ValueError as exc:messagebox.showerror(tr('链接无效'),tr(str(exc)));return
        webbrowser.open(url)

    def use_contact(self):
        if self.busy or self.current is None:return
        selection=self.contact_tree.selection()
        if not selection:return
        candidate_index=int(selection[0])
        self.save_current(refresh=False)
        row=self.rows[self.current]
        candidate=json.loads(row['email_candidates'])[candidate_index]
        row.update(researcher_name=candidate['name'],researcher_email=candidate['email'],
                   email_source_url=candidate['source_url'],contact_basis=candidate['basis'],reviewed='')
        for k,v in self.vars.items():v.set(row.get(k,''))
        self.reviewed.set(False);self.dirty=False
        save_table(DATA/'mailing_list.csv',self.fields,self.rows)
        self.refresh();self.render_preview()

    def find_emails(self):
        if self.busy:return
        self.save_current(refresh=False)
        selected=[int(i) for i in self.tree.selection()]
        indices=selected or [i for i,r in enumerate(self.rows) if not r.get('researcher_email') and r.get('already_in_relations','').lower() not in ('true','1','yes')]
        if not indices:messagebox.showinfo(tr('无需查找'),tr('没有选中记录，也没有缺少邮箱的记录。'));return
        if len(indices)>20 and not messagebox.askyesno(tr('批量读取公开网页'),tr(f'将查找 {len(indices)} 条记录。软件会打开浏览器读取项目和个人主页，并缓存重复页面。\n可随时点击“停止查找”，已完成的结果会保留。继续？')):return
        self.lookup_stop.clear();self.busy=True
        for w in self.edit_controls:w.configure(state='disabled')
        self.stop_button.config(state='normal')
        for b in self.action_buttons:b.config(state='disabled')
        self.status.set('正在打开浏览器查找公开邮箱… 无需邮箱登录，不会发送邮件。')
        jobs=[(i,dict(self.rows[i])) for i in indices]
        threading.Thread(target=self.lookup_worker,args=(jobs,),daemon=True).start()

    def lookup_worker(self,jobs):
        done=0;failed=0
        log_path=DATA/'email_lookup.log'
        def report(text):
            with log_path.open('a',encoding='utf-8') as log:log.write(text+'\n')
            self.events.put(('lookup_progress',text))
        log_path.write_text('邮箱查找开始\n',encoding='utf-8')
        try:
            with email_lookup.BrowserFetcher(self.lookup_stop,report) as fetch:
                for index,row in jobs:
                    if self.lookup_stop.is_set():break
                    self.events.put(('lookup_progress',f'正在查找 {done+1}/{len(jobs)}：{row.get("project_code", "")}'))
                    try:
                        result=email_lookup.lookup(row,fetch)
                    except (ValueError,email_lookup.PageUnavailable) as exc:
                        result={'email_lookup_status':'此条查找失败，可重试','email_lookup_warnings':str(exc)}
                        report(str(exc));failed+=1
                    self.events.put(('lookup_row',(index,result)))
                    done+=1
            self.events.put(('ok',f'邮箱查找{"已停止" if self.lookup_stop.is_set() else "完成"}：处理 {done} 条，其中 {failed} 条读取失败。请查看“邮箱来源与候选”，核对后再发送。'))
        except email_lookup.LookupStopped:self.events.put(('ok',f'邮箱查找已停止：已完成 {done} 条，结果已保留。'))
        except Exception as exc:
            with log_path.open('a',encoding='utf-8') as log:log.write(traceback.format_exc())
            self.events.put(('error',f'邮箱查找中断，已完成 {done} 条并保留结果。\n{exc}\n\n详细记录：{log_path}'))

    def import_file(self):
        if self.busy:return
        path=filedialog.askopenfilename(filetypes=[('CSV / Excel','*.csv *.xlsx')])
        if not path:return
        self.save_current(refresh=False)
        if self.rows and not messagebox.askyesno(tr('导入新的清单'),tr('新清单会替换当前工作清单，发送历史会保留。是否继续？\n需要保留旧清单时请先导出。')):return
        self.open_path(path)

    def open_path(self,path):
        try:
            fields,rows,info=load_table(path)
            save_table(DATA/'mailing_list.csv',fields,rows)
            self.fields,self.rows,self.current,self.dirty=fields,rows,None,False
            for v in self.vars.values():v.set('')
            self.dirty=False;self.refresh()
            self.status.set(f'{len(rows):,} 条记录 · {info} · 原文件未修改')
        except Exception as exc:messagebox.showerror(tr('无法导入'),tr(str(exc)))

    def export(self):
        if self.busy:return
        self.save_current(refresh=False)
        path=filedialog.asksaveasfilename(defaultextension='.csv',initialfile='personal_mailing_list.csv',filetypes=[('CSV UTF-8','*.csv')])
        if path:
            save_table(path,self.fields,self.rows);self.status.set('已导出清单（CSV UTF-8）')

    def account_settings(self):
        if self.busy:return
        win=tk.Toplevel(self);win.title(tr('个人邮箱设置'));win.geometry(f'{self.px(680)}x{self.px(700)}');win.transient(self);win.grab_set()
        frame=ttk.Frame(win,padding=20);frame.pack(fill='both',expand=True)
        ttk.Label(frame,text=tr('选择发信方式并配置已获准使用的个人邮箱'),font=(UI_FONT,13,'bold')).pack(anchor='w',pady=(0,10))
        ttk.Label(frame,text=tr('Windows 可使用经典版 Outlook；macOS 和 Windows 均可使用 SMTP。软件不保存邮箱密码。'),wraplength=600).pack(anchor='w',pady=(0,10))
        ttk.Label(frame,text=tr('发信方式')).pack(anchor='w')
        methods={'Classic Outlook (Windows)':'outlook','SMTP (macOS / Windows)':'smtp'}
        choices=list(methods) if sys.platform=='win32' else ['SMTP (macOS / Windows)']
        method=tk.StringVar(value=next(label for label,value in methods.items() if value==self.settings.get('transport','smtp')))
        method_select=ttk.Combobox(frame,textvariable=method,values=choices,state='readonly')
        method_select.pack(fill='x',pady=(4,10))
        values={}
        for key,label in [('sender_name','你的姓名（英文署名）'),('sender_email','个人发件邮箱'),('reply_to','回复邮箱（可留空）')]:
            ttk.Label(frame,text=tr(label)).pack(anchor='w')
            values[key]=tk.StringVar(value=self.settings.get(key,''));ttk.Entry(frame,textvariable=values[key]).pack(fill='x',pady=(4,10))
        smtp_frame=ttk.Frame(frame)
        for key,label in [('smtp_host','SMTP 服务器'),('smtp_port','SMTP 端口'),('smtp_username','SMTP 登录邮箱')]:
            ttk.Label(smtp_frame,text=tr(label)).pack(anchor='w')
            values[key]=tk.StringVar(value=self.settings.get(key,''));ttk.Entry(smtp_frame,textvariable=values[key]).pack(fill='x',pady=(3,7))
        ttk.Label(smtp_frame,text=tr('连接安全')).pack(anchor='w')
        security=tk.StringVar(value=self.settings.get('smtp_security','starttls'))
        ttk.Combobox(smtp_frame,textvariable=security,values=('starttls','ssl'),state='readonly').pack(fill='x',pady=(3,7))
        ttk.Label(smtp_frame,text=tr('登录方式')).pack(anchor='w')
        auth_methods={'Microsoft sign-in':'microsoft','Password environment variable':'password'}
        auth=tk.StringVar(value=next(label for label,value in auth_methods.items() if value==self.settings.get('auth_mode','microsoft')))
        auth_select=ttk.Combobox(smtp_frame,textvariable=auth,values=list(auth_methods),state='readonly')
        auth_select.pack(fill='x',pady=(3,7))
        auth_frame=ttk.Frame(smtp_frame)
        auth_frame.pack(fill='x')
        auth_entries={}
        for key,label in [('tenant_id','Microsoft 租户 ID'),('client_id','Microsoft 应用 ID'),('password_env','密码环境变量名称')]:
            block=ttk.Frame(auth_frame)
            ttk.Label(block,text=tr(label)).pack(anchor='w')
            values[key]=tk.StringVar(value=self.settings.get(key,''))
            ttk.Entry(block,textvariable=values[key]).pack(fill='x',pady=(3,7))
            auth_entries[key]=block
        def show_auth(*_):
            for block in auth_entries.values():block.pack_forget()
            keys=('tenant_id','client_id') if auth_methods[auth.get()]=='microsoft' else ('password_env',)
            for key in keys:auth_entries[key].pack(fill='x')
        def show_smtp(*_):
            if methods[method.get()]=='smtp':smtp_frame.pack(fill='x',pady=(0,8),before=save_button)
            else:smtp_frame.pack_forget()
        auth_select.bind('<<ComboboxSelected>>',show_auth)
        method_select.bind('<<ComboboxSelected>>',show_smtp)
        def save():
            candidate=dict(self.settings);candidate.update({k:v.get().strip() for k,v in values.items()})
            candidate.update(transport=methods[method.get()],smtp_security=security.get(),auth_mode=auth_methods[auth.get()])
            if candidate['transport']=='smtp' and not candidate['smtp_username']:
                candidate['smtp_username']=candidate['sender_email']
            try:
                core.validate_config(candidate)
                if candidate['transport'] == 'smtp':core.validate_smtp_port(candidate)
            except Exception as exc:messagebox.showerror(tr('请检查填写内容'),tr(str(exc)),parent=win);return
            self.settings=candidate
            (DATA/'settings.json').write_text(json.dumps(candidate,ensure_ascii=False,indent=2),encoding='utf-8')
            self.update_account_label();self.render_preview();win.destroy()
        save_button=ttk.Button(frame,text=tr('保存个人邮箱'),command=save,style='Primary.TButton')
        save_button.pack(anchor='e')
        show_auth();show_smtp()

    def check_connection(self):
        if self.busy:return
        try:core.validate_config(self.settings,True)
        except Exception as exc:messagebox.showerror(tr('先设置个人邮箱'),tr(str(exc)));return
        self.start_worker('check',[],dict(self.settings))

    def dispatch(self,command):
        if self.busy:return
        if command=='draft' and self.settings.get('transport')!='outlook':
            messagebox.showinfo(tr('生成草稿'),tr('SMTP 不支持 Outlook 草稿；可以先试发给自己。'))
            return
        self.save_current(refresh=False)
        selected=list(self.tree.selection())
        if not selected:messagebox.showinfo(tr('选择记录'),tr('请先选择左侧需要处理的记录。'));return
        try:core.validate_config(self.settings,True)
        except Exception as exc:messagebox.showerror(tr('先设置个人邮箱'),tr(str(exc)));return
        self.read_history()
        rows=[];keys=set()
        for index in selected:
            row=self.rows[int(index)]
            if core.validate(row):continue
            key=core.key_for(row)
            if key in keys or (command!='test' and key in self.history):continue
            keys.add(key);rows.append(dict(row))
        if not rows:messagebox.showinfo(tr('没有可处理的记录'),tr('请补齐信息并勾选已核对。已生成草稿或已提交的记录不会重复发送。'));return
        if command=='test':rows=rows[:1]
        if command=='send':
            recipients='\n'.join(r['researcher_email'] for r in rows[:6])
            if len(rows)>6:recipients+=f'\n……共 {len(rows)} 封'
            if not messagebox.askyesno(tr('发送个人确认邮件'),tr(f'将从 {self.settings["sender_email"]} 发送 {len(rows)} 封邮件。\n\n{recipients}\n\n请先在右侧预览内容。现在发送？')):return
        elif command=='draft':
            if not messagebox.askyesno(tr('保存到 Outlook 草稿箱'),tr(f'将生成 {len(rows)} 封草稿，不会发送。\n草稿随后需在 Outlook 中手动发送；软件不会再自动发送这些记录。\n是否继续？')):return
        self.start_worker(command,rows,dict(self.settings))

    def start_worker(self,command,rows,config):
        self.busy=True
        self.auth_stop.clear()
        self.auth_pending.clear()
        for w in self.edit_controls:w.configure(state='disabled')
        self.status.set('正在处理… 请保持应用打开，等待结果。')
        for b in self.action_buttons:b.config(state='disabled')
        threading.Thread(target=self.worker,args=(command,rows,config),daemon=True).start()

    def auth_prompt(self,uri,code):
        self.auth_pending.set()
        self.events.put(('auth',(uri,code)))

    def auth_done(self):
        self.auth_pending.clear()
        self.events.put(('auth_done',None))

    def cancel_sign_in(self):
        if self.auth_pending.is_set():
            self.auth_stop.set()
        self.finish_auth()
        if self.auth_stop.is_set():
            self.status.set('正在取消 Microsoft 登录…')

    def finish_auth(self):
        if self.auth_window is not None:
            self.auth_window.destroy()
            self.auth_window = None
        self.stop_button.config(text=tr('停止查找'),command=self.lookup_stop.set,state='disabled')

    def show_auth_prompt(self,uri,code):
        self.stop_button.config(text=tr('取消登录'),command=self.cancel_sign_in,state='normal')
        self.status.set('等待 Microsoft 登录；可点击“取消登录”。')
        win=tk.Toplevel(self)
        win.title(tr('Microsoft 登录'))
        win.transient(self)
        frame=ttk.Frame(win,padding=20);frame.pack(fill='both',expand=True)
        ttk.Label(frame,text=tr(f'请在打开的网页输入代码 {code} 完成登录。\n\n登录网址：{uri}'),wraplength=420).pack(pady=(0,12))
        ttk.Button(frame,text=tr('取消登录'),command=self.cancel_sign_in).pack(anchor='e')
        win.protocol('WM_DELETE_WINDOW',self.cancel_sign_in)
        self.auth_window=win
        webbrowser.open(uri)

    def worker(self,command,rows,config):
        com=None
        try:
            if config.get('transport')=='outlook':
                import pythoncom
                com=pythoncom;com.CoInitialize()
            if command=='check':
                with core.connect(config,on_auth=self.auth_prompt,stop_event=self.auth_stop,on_auth_done=self.auth_done):pass
                result='已连接经典版 Outlook，并找到指定个人邮箱。未创建或发送邮件。' if config.get('transport')=='outlook' else 'SMTP 登录与连接成功。未发送邮件。'
                self.events.put(('ok',result))
                return
            save_table(DATA/'batch.csv',self.fields,rows)
            config_path=DATA/'batch_config.json'
            config_path.write_text(json.dumps(config),encoding='utf-8')
            args=SimpleNamespace(command=command,input=DATA/'batch.csv',config=config_path,output=DATA/'validation',limit=len(rows),to=config['sender_email'])
            log=io.StringIO()
            with redirect_stdout(log):core.run(args,on_auth=self.auth_prompt,stop_event=self.auth_stop,on_auth_done=self.auth_done)
            (DATA/'last_operation.txt').write_text(log.getvalue(),encoding='utf-8')
            text={'test':'已将 1 封测试邮件交给发件服务，收件人为你自己的邮箱。',
                  'draft':'草稿处理完成，请在 Outlook 草稿箱中检查并手动发送。',
                  'send':'选中邮件处理完成，请在邮箱中检查已发送、退信和回复。'}[command]
            self.events.put(('ok',text))
        except core.SignInCancelled:self.events.put(('cancelled','Microsoft 登录已取消，未发送邮件。'))
        except Exception as exc:self.events.put(('error',str(exc)))
        finally:
            self.auth_pending.clear()
            if com:com.CoUninitialize()

    def poll(self):
        try:
            kind,text=self.events.get_nowait()
            if kind=='auth':
                if self.auth_pending.is_set() and not self.auth_stop.is_set():
                    uri,code=text
                    self.show_auth_prompt(uri,code)
                self.after(100,self.poll);return
            if kind=='auth_done':
                self.finish_auth()
                self.after(100,self.poll);return
            if kind=='lookup_progress':
                self.status.set(text)
                self.after(100,self.poll);return
            if kind=='lookup_row':
                index,result=text
                self.rows[index]=email_lookup.merge_result(self.rows[index],result)
                save_table(DATA/'mailing_list.csv',self.fields,self.rows)
                if self.current==index:
                    for k,v in self.vars.items():v.set(self.rows[index].get(k,''))
                    self.reviewed.set(self.rows[index].get('reviewed','').lower()=='yes');self.dirty=False
                    self.render_preview()
                self.after(50,self.poll);return
            self.busy=False
            for w in self.edit_controls:w.configure(state='normal')
            self.finish_auth()
            for b in self.action_buttons:b.config(state='normal')
            self.update_account_label()
            if kind=='cancelled' and self.close_when_idle:
                self.close_when_idle=False
                self.save_current(refresh=False)
                self.destroy()
                return
            self.close_when_idle=False
            self.refresh();self.render_preview();self.status.set(text)
            if kind=='error':messagebox.showerror(tr('操作未完成'),tr(text+'\n\n若已开始处理邮件，请先检查邮箱的已发送、发件箱和退信，再决定是否重试。'))
            elif kind!='cancelled':messagebox.showinfo(tr('处理结果'),tr(text))
        except queue.Empty:pass
        self.after(150,self.poll)

    def close_app(self):
        if self.busy:
            if self.auth_pending.is_set() or self.auth_stop.is_set():
                self.close_when_idle=True
                self.cancel_sign_in()
                return
            messagebox.showinfo(tr('正在处理'),tr('请等待本批操作结束后关闭，避免发送结果不确定。'));return
        self.save_current(refresh=False);self.destroy()

if __name__=='__main__':
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:pass
    app=App()
    if '--self-check' in sys.argv:
        app.withdraw()
        def self_check():
            try:
                import openpyxl, msal
                from playwright.sync_api import sync_playwright
                if sys.platform == 'win32':
                    import win32com.client, pythoncom
                with sync_playwright():pass
                with email_lookup.BrowserFetcher() as browser:
                    browser_version=browser.browser.version
                report={'startup':'ok','rows':len(app.rows),'widgets':len(app.vars),'excel':'available',
                        'smtp_adapter':'available','outlook_adapter':'available' if sys.platform=='win32' else 'not applicable',
                        'browser_driver':'available','browser_version':browser_version,'emails_sent':0}
                Path(sys.argv[sys.argv.index('--self-check')+1]).write_text(json.dumps(report),encoding='utf-8')
            finally:app.destroy()
        app.after(200,self_check)
    app.mainloop()
